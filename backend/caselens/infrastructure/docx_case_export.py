"""Builds the full decision(s) as a Word file: everything the Court printed, nothing reworded."""
import io
from datetime import date

import docx
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor

from caselens.application.ports.gateways import CaseDocumentExporter
from caselens.domain.entities import Case, Footnote
from caselens.domain.services.decision_blocks import BlockKind, DecisionBlock, DecisionBlocks, MarkerPiece

_GREY = RGBColor(0x59, 0x59, 0x59)
_FONT = "Times New Roman"
_COVER_TITLE_CHARS = 110  # a petitioners' list can run to hundreds of names; the full caption is in the decision itself
_NO_FOOTNOTE_TEXT = (
    "This decision has footnote numbers, but their text could not be read from the Court's page. "
    "See the footnotes on the official page."
)
_NOTICE = (
    "Court text from Lawphil (Arellano Law Foundation), https://lawphil.net. Informational only, not legal advice; "
    "Lawphil gives no warranty of accuracy or completeness: confirm with the Supreme Court."
)


class DocxCaseExporter(CaseDocumentExporter):
    def __init__(self, blocks: DecisionBlocks | None = None) -> None:
        self._blocks = blocks or DecisionBlocks()

    def export(self, cases: list[Case]) -> bytes:
        document = docx.Document()
        normal = document.styles["Normal"]
        normal.font.name, normal.font.size = _FONT, Pt(12)
        if len(cases) > 1:
            self._cover(document, cases)
        for number, case in enumerate(cases):
            if number > 0 or len(cases) > 1:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            self._case(document, case)
        out = io.BytesIO()
        document.save(out)
        return out.getvalue()

    # -- parts -------------------------------------------------------------------------

    def _cover(self, document: DocxDocument, cases: list[Case]) -> None:
        document.add_heading("Full cases", level=1)
        self._grey(document, f"Made by CaseLens on {date.today():%B %d, %Y}. {len(cases)} cases.")
        for case in cases:
            document.add_paragraph(f"{self._numbers(case)}: {self._cover_title(case)}", style="List Bullet")
        self._grey(document, _NOTICE, size=10)

    def _case(self, document: DocxDocument, case: Case) -> None:
        document.add_heading(self._numbers(case), level=1)  # the Court's own full caption follows in the decision
        meta = [
            f"{case.decision_date:%B %d, %Y}" if case.decision_date else None,
            case.division,
            f"Written by Justice {case.ponente}" if case.ponente else None,
        ]
        self._grey(document, " · ".join(part for part in meta if part))
        self._grey(document, f"Official page: {case.source_url}")
        self._grey(document, _NOTICE, size=10)

        self._decision(document, case.full_text)
        self._footnotes(document, case.footnotes, case.full_text)
        for opinion in case.opinions:
            kind = {"dissenting": "Dissenting opinion", "concurring": "Concurring opinion"}.get(opinion.kind, "Opinion")
            document.add_heading(f"{kind}, Justice {opinion.author}" if opinion.author else kind, level=2)
            self._decision(document, opinion.text)
            self._footnotes(document, opinion.footnotes, opinion.text)

    def _decision(self, document: DocxDocument, text: str) -> None:
        for block in self._blocks.split(text):
            self._block(document, block)

    def _block(self, document: DocxDocument, block: DecisionBlock) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        centred = block.kind in (BlockKind.CAPTION, BlockKind.PARTIES, BlockKind.TITLE, BlockKind.NOTICE)
        if centred:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        bold = block.kind in (BlockKind.PARTIES, BlockKind.TITLE, BlockKind.PONENTE, BlockKind.HEADING, BlockKind.NOTICE)
        for piece in self._blocks.pieces(block.text):
            if isinstance(piece, MarkerPiece):
                run = paragraph.add_run(str(piece.number))
                run.font.superscript = True
            else:
                run = paragraph.add_run(piece.text)
            run.bold = bold
            if block.kind is BlockKind.TITLE:
                run.font.size = Pt(14)

    def _footnotes(self, document: DocxDocument, footnotes: list[Footnote], text: str) -> None:
        if not footnotes:
            if "[^" in text:
                self._grey(document, _NO_FOOTNOTE_TEXT, size=10)
            return
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(12)
        heading.add_run("Footnotes").bold = True
        for footnote in footnotes:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(f"{footnote.number}. {footnote.text}" if footnote.text else f"{footnote.number}.")
            run.font.size = Pt(10)

    @staticmethod
    def _grey(document: DocxDocument, text: str, *, size: float | None = None) -> None:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        run.font.color.rgb = _GREY
        if size:
            run.font.size = Pt(size)

    @staticmethod
    def _cover_title(case: Case) -> str:
        title = " ".join((case.title or "").split())
        return title if len(title) <= _COVER_TITLE_CHARS else title[: _COVER_TITLE_CHARS - 1].rstrip(" ,") + "…"

    @staticmethod
    def _numbers(case: Case) -> str:
        numbers = case.all_numbers
        return f"G.R. Nos. {', '.join(numbers)}" if len(numbers) > 1 else f"G.R. No. {numbers[0]}"
