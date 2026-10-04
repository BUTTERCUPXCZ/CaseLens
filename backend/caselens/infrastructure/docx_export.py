"""Builds the finished reviewer as a Word file: the student's reviewer with a boxed digest after each citing paragraph."""
import io

import docx
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

from caselens.application.ports.gateways import ReviewerDocumentExporter
from caselens.domain.finished_reviewer import FinishedReviewer, ReviewerBox

_AFTER_BORDERS = {qn(f"w:{name}") for name in ("shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription")}
_GREY = RGBColor(0x59, 0x59, 0x59)
_ATTRIBUTION = (
    "Court text from Lawphil (Arellano Law Foundation), https://lawphil.net. Informational only, not legal advice; "
    "Lawphil gives no warranty of accuracy or completeness: confirm with the Supreme Court. "
    "Explanations marked 'Drafted from the decision' were written by an AI service from that decision: check them."
)
_PDF_NOTE = "Rebuilt from your PDF: the text is the same, but the fonts and layout differ from the original."


class DocxDigestExporter(ReviewerDocumentExporter):
    def export(self, reviewer: FinishedReviewer, original: bytes | None) -> bytes:
        if reviewer.source == "docx" and original is not None:
            document = docx.Document(io.BytesIO(original))
            self._insert_into_copy(document, reviewer)
        else:
            document = self._rebuild(reviewer)
        self._append_unplaced(document, reviewer)
        self._attribution(document)
        out = io.BytesIO()
        document.save(out)
        return out.getvalue()

    # -- the two ways to lay it out --------------------------------------------------

    def _insert_into_copy(self, document: DocxDocument, reviewer: FinishedReviewer) -> None:
        """The student's own file, untouched except for the boxes inserted after the citing paragraphs."""
        paragraphs = document.paragraphs
        for index, boxes in reviewer.after_block.items():
            anchor = paragraphs[index]._p
            for box in boxes:
                for element in self._box_elements(document, box):
                    anchor.addnext(element)
                    anchor = element

    def _rebuild(self, reviewer: FinishedReviewer) -> DocxDocument:
        document = docx.Document()
        document.add_heading(f"{reviewer.filename} (with digests)", level=1)
        if reviewer.source == "pdf":
            self._plain(document, _PDF_NOTE, italic=True)
        for index, text in enumerate(reviewer.blocks):
            level = reviewer.heading_levels[index] if index < len(reviewer.heading_levels) else None
            if level is None:
                document.add_paragraph(text)
            else:
                document.add_heading(text, level=level + 1)  # level 1 is the document title (Heading 1 is used for the file name)
            anchor = document.paragraphs[-1]._p
            for box in reviewer.after_block.get(index, []):
                for element in self._box_elements(document, box):
                    anchor.addnext(element)
                    anchor = element
        return document

    def _append_unplaced(self, document: DocxDocument, reviewer: FinishedReviewer) -> None:
        if not reviewer.unplaced:
            return
        document.add_heading("Cases cited that we could not place next to their paragraph", level=2)
        for box in reviewer.unplaced:
            self._box_elements(document, box, at_end=True)  # left where `add_*` put them: at the end of the file

    def _attribution(self, document: DocxDocument) -> None:
        self._plain(document, _ATTRIBUTION, italic=True, size=8)

    # -- one digest box ---------------------------------------------------------------

    def _box_elements(self, document: DocxDocument, box: ReviewerBox, at_end: bool = False) -> list:
        """The box as XML elements: a title line, the bordered table, and a spacer paragraph. They are created at the end of
        the document and (unless `at_end`) detached so the caller can place them after the right paragraph."""
        title = document.add_paragraph()
        run = title.add_run(box.title)
        run.bold = True
        run.font.size = Pt(15)

        table = document.add_table(rows=1, cols=1)
        try:
            table.style = "Table Grid"  # present in Word's default template; the explicit borders below cover files without it
        except KeyError:
            pass
        self._borders(table)
        cell = table.rows[0].cells[0]
        first = cell.paragraphs[0]
        header = first.add_run(box.heading)
        header.bold = True
        header.font.size = Pt(13)
        self._line(cell, f"Court's record: {box.source_url}", size=8, colour=_GREY)
        if box.check_note:
            self._line(cell, box.check_note, italic=True, size=9)
        if not box.ready:
            self._line(cell, "Some explanations are still being written. Download again in a minute.", italic=True, size=9, colour=_GREY)
        for item in box.fields:
            self._line(cell, item.label, bold=True, size=14, space_before=12)
            for part in [p for p in item.text.split("\n\n") if p.strip()] or [""]:
                self._line(cell, part, size=11)
            if item.note:
                self._line(cell, item.note, italic=True, size=8, colour=_GREY)
        spacer = document.add_paragraph()

        elements = [title._p, table._tbl, spacer._p]
        if not at_end:
            for element in elements:
                element.getparent().remove(element)
        return elements

    @staticmethod
    def _line(cell, text: str, *, bold=False, italic=False, size: float | None = None, colour=None, space_before=0) -> Paragraph:
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.space_before = Pt(space_before)
        run = paragraph.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if colour is not None:
            run.font.color.rgb = colour
        return paragraph

    @staticmethod
    def _plain(document: DocxDocument, text: str, *, italic=False, size: float | None = None) -> None:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(text)
        run.italic = italic
        run.font.color.rgb = _GREY
        if size:
            run.font.size = Pt(size)

    @staticmethod
    def _borders(table) -> None:
        """A thin border on every side, set directly so it works even if the student's file has no 'Table Grid' style."""
        properties = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right"):
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), "6")
            edge.set(qn("w:space"), "0")
            edge.set(qn("w:color"), "444444")
            borders.append(edge)
        # Word's schema fixes the order of tblPr children: borders come before shading, layout, margins and tblLook.
        anchor = next((child for child in properties if child.tag in _AFTER_BORDERS), None)
        if anchor is not None:
            anchor.addprevious(borders)
        else:
            properties.append(borders)
