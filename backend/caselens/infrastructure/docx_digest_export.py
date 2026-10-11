"""The case digest as a Word file, laid out like the client's sample: "CASE DIGEST", the case name, the citation line,
the topic and ponente, then one heading per section. "Case Summary and Doctrine" follows the client's reference instead: one top line
("Nobleza v. Nuega, G.R. No. 193038 (March 11, 2015), 752 SCRA 602"), the summary's paragraphs, then the numbered doctrines, each
one paragraph led by its label in bold."""
import io
from datetime import date

import docx
from docx.document import Document as DocxDocument
from docx.shared import Pt, RGBColor

from caselens.domain.digest_v2 import SECTION_TITLES, DigestBlock, DigestDraft, DigestHeader, Level, Section, block_heading, sections_for
from caselens.infrastructure.source_notice import digest_notice

_GREY = RGBColor(0x59, 0x59, 0x59)
_REMEMBER = "Remember for recitation:"
_SUBSECTIONS = {Section.ARGUMENTS_PETITIONERS, Section.ARGUMENTS_RESPONDENTS}  # printed as a Heading 2 under Facts


class DocxCaseDigestExporter:
    def export(self, header: DigestHeader, draft: DigestDraft, level: Level, *, show_sources: bool = False, source_url: str | None = None) -> bytes:
        document = docx.Document()
        document.styles["Normal"].font.name = "Calibri"
        document.styles["Normal"].font.size = Pt(11)
        if level is Level.SHORT:
            self._one_line(document, header)
        else:
            self._header(document, header)
        for section in sections_for(level, draft):  # in the level's order: the short file is Case Summary, then Doctrine
            if section not in draft.sections:
                continue
            title = SECTION_TITLES[section]
            document.add_heading(title, level=2 if section in _SUBSECTIONS else 1)
            blocks = draft.sections[section]
            if section is Section.WHY and blocks and blocks[0].as_list:
                document.add_paragraph().add_run(_REMEMBER).bold = True  # the sample's lead-in to the recitation points
            for block in blocks:
                block = DigestBlock(block.sentences, block_heading(section, block.heading), block.as_list)  # no "Doctrine" under Doctrine
                if section is Section.DOCTRINE and not block.as_list:
                    self._doctrine(document, block, show_sources)
                else:
                    self._block(document, block, show_sources, name_first=section is Section.DISSENTS, plain=section is Section.CASE_SUMMARY)
        note = document.add_paragraph()
        run = note.add_run(digest_notice(source_url))
        run.italic, run.font.size, run.font.color.rgb = True, Pt(9), _GREY
        out = io.BytesIO()
        document.save(out)
        return out.getvalue()

    @staticmethod
    def _one_line(document: DocxDocument, header: DigestHeader) -> None:
        """"Nobleza v. Nuega, G.R. No. 193038 (March 11, 2015), 752 SCRA 602", the case name in italics, as the reference prints it."""
        line = document.add_paragraph()
        name = line.add_run(header.case_name)
        name.italic = True
        rest = header.one_line[len(header.case_name):]
        if rest:
            line.add_run(rest)

    @classmethod
    def _doctrine(cls, document: DocxDocument, block: DigestBlock, show_sources: bool) -> None:
        """One doctrine as one paragraph: "1. Innocent purchaser for value: burden and due diligence:" in bold, then its sentences."""
        paragraph = document.add_paragraph()
        if block.heading:
            paragraph.add_run(block.heading.rstrip(":") + ": ").bold = True
        paragraph.add_run(" ".join(sentence.text for sentence in block.sentences))
        cls._sources(document, block, show_sources)

    @staticmethod
    def _header(document: DocxDocument, header: DigestHeader) -> None:
        document.add_paragraph().add_run("CASE DIGEST").bold = True
        title = document.add_paragraph()
        run = title.add_run(header.case_name)
        run.bold, run.font.size = True, Pt(16)
        document.add_paragraph(header.citation)
        facts = [f"Topic: {header.topic}" if header.topic else None, f"Ponente: {header.ponente}" if header.ponente else None]
        line = "  |  ".join(part for part in facts if part)
        if line:
            document.add_paragraph(line)

    @classmethod
    def _block(cls, document: DocxDocument, block: DigestBlock, show_sources: bool, *, name_first: bool = False, plain: bool = False) -> None:
        """`plain`: no bold key sentences (the Case Summary, as the reference prints it)."""
        if block.heading:
            document.add_heading(block.heading, level=2)
        if block.as_list:
            for sentence in block.sentences:
                item = document.add_paragraph(style="List Bullet")
                name, sep, rest = sentence.text.partition(": ")
                if name_first and sep and len(name) <= 60:  # "Cruz, J.: ..." the justice's name in bold, as the sample prints the dissents
                    item.add_run(name + ":").bold = True
                    item.add_run(" " + rest).bold = sentence.key
                else:
                    item.add_run(sentence.text).bold = sentence.key
        elif block.sentences:
            paragraph = document.add_paragraph()
            for n, sentence in enumerate(block.sentences):
                paragraph.add_run(("" if n == 0 else " ") + sentence.text).bold = sentence.key and not plain  # a key sentence is bold, as in the sample
        cls._sources(document, block, show_sources)

    @staticmethod
    def _sources(document: DocxDocument, block: DigestBlock, show_sources: bool) -> None:
        if show_sources:
            numbers = list(dict.fromkeys(c[1:] for s in block.sentences for c in s.cites if c.startswith("P")))
            if numbers:
                shown = ", ".join(numbers[:12]) + (" and more" if len(numbers) > 12 else "")
                line = document.add_paragraph()
                run = line.add_run(f"Based on decision paragraphs {shown}.")
                run.font.size, run.font.color.rgb = Pt(8.5), _GREY
