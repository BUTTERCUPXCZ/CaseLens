import io

import docx

from caselens.application.ports.gateways import DocumentTextExtractor
from caselens.domain.errors import DocumentExtractionError


class DocxTextExtractor(DocumentTextExtractor):
    """Reads paragraphs and table cells from .docx files."""

    def supports(self, filename: str) -> bool:
        return filename.lower().endswith(".docx")

    def extract(self, data: bytes) -> str:
        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise DocumentExtractionError(f"Cannot read DOCX: {exc}") from exc

        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)

        text = "\n".join(parts)
        if not text.strip():
            raise DocumentExtractionError("DOCX has no text.")
        return text
