"""The Word file: the student's reviewer with a bordered digest box after each citing paragraph."""
import io
from datetime import date

import docx
import pytest

from caselens.application.use_cases.build_finished_reviewer import BuildFinishedReviewer
from caselens.domain.finished_reviewer import BoxField, FinishedReviewer, ReviewerBox
from caselens.infrastructure.docx_export import DocxDigestExporter
from caselens.infrastructure.extraction.block_reader import CompositeBlockReader, DocxBlockReader, PdfBlockReader
from tests.helpers import FIXTURES, REVIEWER_PARAGRAPHS, reviewer_docx


def box(number=1, ready=True, check_note=None, fields=None):
    return ReviewerBox(
        digest_id=number, case_id=number, citation_id=number, title=f"Digest {number}: Facts, Issue, Ruling and Doctrine",
        heading="Review Center Associations of the Philippines vs. Executive Secretary Eduardo Ermita, et al., G.R. No. 180046 (April 2, 2009)",
        source_url="https://lawphil.net/judjuris/juri2009/apr2009/gr_180046_2009.html", decided_on=date(2009, 4, 2),
        check_note=check_note, ready=ready,
        fields=fields or (
            BoxField("Facts", "On 11 and 12 June 2006, the PRC conducted the Nursing Board Examinations."),
            BoxField("Issue", ""),
            BoxField("Topic explained", "The case tests whether EO 566 is valid.", "Drafted from the decision (decision paragraphs 7, 65). Check it."),
            BoxField("Why this case matters", "", "Still being written."),
        ),
    )


def body_order(document):
    """The document body as a list of 'p:<text>' / 'table' in order."""
    items = []
    for element in document.element.body:
        tag = element.tag.split("}")[1]
        if tag == "p":
            items.append("p:" + "".join(t.text or "" for t in element.iter() if t.tag.endswith("}t")))
        elif tag == "tbl":
            items.append("table")
    return items


def export_docx(boxes_after=None, unplaced=None, source="docx", ready=True):
    data = reviewer_docx()
    reviewer = FinishedReviewer(1, "reviewer.docx", source, list(REVIEWER_PARAGRAPHS))
    reviewer.after_block = boxes_after if boxes_after is not None else {2: [box(1, ready=ready)]}
    reviewer.unplaced = unplaced or []
    out = DocxDigestExporter().export(reviewer, data)
    return docx.Document(io.BytesIO(out))


def test_the_box_goes_right_after_the_citing_paragraph_and_the_students_text_is_untouched():
    document = export_docx()
    order = body_order(document)
    assert order[:3] == ["p:" + REVIEWER_PARAGRAPHS[0], "p:" + REVIEWER_PARAGRAPHS[1], "p:" + REVIEWER_PARAGRAPHS[2]]
    assert order[3].startswith("p:Digest 1: Facts, Issue, Ruling and Doctrine") and order[4] == "table"
    assert order[5] == "p:" and order[6] == "p:" + REVIEWER_PARAGRAPHS[3]  # a spacer, then the student's next paragraph
    originals = [p.text for p in document.paragraphs if p.text in REVIEWER_PARAGRAPHS]
    assert originals == REVIEWER_PARAGRAPHS  # every original paragraph is still there, in order, unchanged


def test_the_box_has_the_courts_header_source_fields_and_where_each_answer_came_from():
    cell = export_docx().tables[0].rows[0].cells[0]
    lines = [p.text for p in cell.paragraphs if p.text.strip() != "" or True]
    text = "\n".join(lines)
    assert "G.R. No. 180046 (April 2, 2009)" in lines[0]
    assert "Court's record: https://lawphil.net/judjuris/juri2009/apr2009/gr_180046_2009.html" in text
    assert "Facts\nOn 11 and 12 June 2006, the PRC conducted the Nursing Board Examinations." in text
    assert "The case tests whether EO 566 is valid.\nDrafted from the decision (decision paragraphs 7, 65). Check it." in text
    assert "Why this case matters\n\nStill being written." in text  # an empty field keeps its label for the student to fill


def test_an_empty_field_keeps_its_label_like_topic_explained_in_the_students_own_boxes():
    text = "\n".join(p.text for p in export_docx().tables[0].rows[0].cells[0].paragraphs)
    assert "Issue\n\n" in text


def test_the_box_has_visible_borders_even_when_the_students_file_has_no_table_style():
    borders = export_docx().tables[0]._tbl.tblPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblBorders"
    )
    assert borders is not None and len(borders) == 4


def test_a_citation_difference_is_pointed_out_in_plain_words():
    reviewer = FinishedReviewer(1, "r.docx", "docx", list(REVIEWER_PARAGRAPHS), {2: [box(check_note="Check this citation. year: your reviewer says 2010; the Court's record says 2009.")]})
    document = docx.Document(io.BytesIO(DocxDigestExporter().export(reviewer, reviewer_docx())))
    assert "your reviewer says 2010; the Court's record says 2009" in "\n".join(p.text for p in document.tables[0].rows[0].cells[0].paragraphs)


def test_a_digest_still_being_written_is_marked_so_the_student_downloads_again():
    text = "\n".join(p.text for p in export_docx(ready=False).tables[0].rows[0].cells[0].paragraphs)
    assert "Download again in a minute" in text


def test_two_boxes_after_the_same_paragraph_keep_their_order():
    document = export_docx({2: [box(1), box(2)]})
    titles = [i for i in body_order(document) if i.startswith("p:Digest")]
    assert [t[:10] for t in titles] == ["p:Digest 1", "p:Digest 2"]
    assert len(document.tables) == 2


def test_a_case_that_could_not_be_placed_is_added_at_the_end_with_its_own_heading():
    document = export_docx({}, unplaced=[box(5)])
    order = body_order(document)
    assert any(item.startswith("p:Cases cited that we could not place") for item in order)
    assert order.index("table") > order.index("p:" + REVIEWER_PARAGRAPHS[3])


def test_the_lawphil_attribution_and_the_ai_notice_close_the_file():
    last = [p.text for p in export_docx().paragraphs if p.text][-1]
    assert "Lawphil (Arellano Law Foundation)" in last and "written by an AI service" in last


def test_a_pdf_is_rebuilt_as_a_new_word_file_and_says_so():
    data = (FIXTURES / "sample_case.pdf").read_bytes()
    blocks = CompositeBlockReader([PdfBlockReader(), DocxBlockReader()]).blocks("sample_case.pdf", data)
    reviewer = FinishedReviewer(1, "sample_case.pdf", "pdf", blocks)
    index = next(i for i, text in enumerate(blocks) if "GR no 180046" in text)
    reviewer.after_block = {index: [box(1)]}

    document = docx.Document(io.BytesIO(DocxDigestExporter().export(reviewer, data)))
    paragraphs = [p.text for p in document.paragraphs]
    assert paragraphs[0] == "sample_case.pdf (with digests)"
    assert "Rebuilt from your PDF" in paragraphs[1]
    assert "PART NINE: LEGISLATIVE DEPARTMENT" in " ".join(paragraphs)  # the reviewer's own text is there
    assert len(document.tables) == 1


def test_the_real_pdf_blocks_are_clean_paragraphs_without_zero_width_marks():
    data = (FIXTURES / "sample_case.pdf").read_bytes()
    blocks = PdfBlockReader().blocks("sample_case.pdf", data)
    assert blocks and all("​" not in b and "\n" not in b for b in blocks)
    assert blocks[0] == "Sample case"
    assert "PART NINE: LEGISLATIVE DEPARTMENT - Article VI 1987 Constitution" in blocks


def test_a_docx_is_read_as_its_own_paragraphs_so_positions_match_the_file():
    blocks = DocxBlockReader().blocks("r.docx", reviewer_docx())
    assert blocks == REVIEWER_PARAGRAPHS
    with pytest.raises(Exception, match="only PDF and Word"):
        CompositeBlockReader([PdfBlockReader(), DocxBlockReader()]).blocks("notes.txt", b"x")


def test_the_border_xml_is_in_the_order_word_requires():
    properties = export_docx().tables[0]._tbl.tblPr
    names = [child.tag.split("}")[1] for child in properties]
    assert names.index("tblBorders") < names.index("tblLook")


def test_the_header_has_no_doubled_full_stop_from_the_courts_caption():
    from caselens.application.use_cases.build_finished_reviewer import _check_note  # noqa: F401  (module imported fine)

    assert "Respondents.," not in "\n".join(p.text for p in export_docx().tables[0].rows[0].cells[0].paragraphs)


def test_the_sample_pdfs_why_this_case_matters_is_a_heading_followed_by_separate_points_not_one_run_together_paragraph():
    blocks = PdfBlockReader().blocks("sample_case.pdf", (FIXTURES / "sample_case.pdf").read_bytes())
    heading = len(blocks) - 1 - blocks[::-1].index("WHY THIS CASE MATTERS")  # the answer's own heading, not the empty template's
    points = blocks[heading + 1 : heading + 4]
    assert [p.split(":")[0] for p in points] == [
        "Upholds the Separation of Powers",
        "Restricts Executive Overreach",
        "Clarifies Administrative Rule-Making",
    ]
    assert all(len(p) < 400 for p in points)  # each point on its own, not the whole answer in one line
    assert "TOPIC EXPLAINED:" in blocks  # the student's other heading is also its own line


def test_a_pdf_without_paragraph_marks_still_splits_headings_and_list_points():
    from caselens.infrastructure.extraction.block_reader import paragraphs_of_block

    block = "Intro line that continues\non the next line.\nTOPIC EXPLAINED:\n- First point here\n- Second point here"
    assert paragraphs_of_block(block, has_paragraph_marks=False) == [
        "Intro line that continues on the next line.",
        "TOPIC EXPLAINED:",
        "- First point here",
        "- Second point here",
    ]


# The order Word's schema requires for the children of a table's properties. Word can refuse or repair a file that breaks it.
_TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
                "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]


def test_every_table_property_is_in_the_order_the_word_schema_requires():
    document = export_docx({2: [box(1), box(2)]})
    for table in document.tables:
        names = [child.tag.split("}")[1] for child in table._tbl.tblPr]
        assert names == sorted(names, key=_TBLPR_ORDER.index)


def test_the_file_is_a_valid_zip_with_well_formed_xml_parts():
    import xml.dom.minidom
    import zipfile

    out = io.BytesIO()
    document = export_docx()
    document.save(out)
    with zipfile.ZipFile(out) as package:
        assert package.testzip() is None
        for name in package.namelist():
            if name.endswith((".xml", ".rels")):
                xml.dom.minidom.parseString(package.read(name))  # raises if the XML is broken
