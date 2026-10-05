"""The full-case Word file, built from REAL saved Lawphil decisions: everything the Court printed, in order,
nothing reworded, footnotes listed, opinions included."""
import io
import re

import docx
import pytest

from caselens.domain.services.decision_blocks import BlockKind, DecisionBlocks
from caselens.infrastructure.docx_case_export import DocxCaseExporter
from caselens.infrastructure.lawphil.html_parser import LawphilCaseParser
from tests.helpers import read_fixture_html

URL = "https://lawphil.net/judjuris/juri2009/apr2009/gr_180046_2009.html"


def parse(name: str, url: str):
    return LawphilCaseParser().parse(read_fixture_html(name), url)


@pytest.fixture(scope="module")
def ermita():
    return parse("gr_180046_2009.html", URL)


def read(data: bytes) -> docx.document.Document:
    return docx.Document(io.BytesIO(data))


def texts(document) -> list[str]:
    return [p.text for p in document.paragraphs]


def test_a_decision_comes_out_in_full_and_in_the_courts_order(ermita):
    document = read(DocxCaseExporter().export([ermita]))
    lines = texts(document)

    court_lines = [re.sub(r"\[\^(\d+)\]", r"\1", line) for line in ermita.full_text.split("\n") if line.strip()]
    # The Court's letter-spaced headings ("D E C I S I O N") are printed closed up, as the on-screen reader does.
    court_lines = [re.sub(r" ", "", line) if re.fullmatch(r"(?:[A-Z] ){3,}[A-Z]", line) else line for line in court_lines]
    decision_part = lines[: lines.index("Footnotes")]  # the opinions come after the decision's own footnotes
    body = [line for line in decision_part if line in court_lines]
    assert body == court_lines  # every line of the decision, none missing, none reordered, none reworded


def test_footnote_markers_become_raised_numbers_and_none_is_left_in_the_text(ermita):
    document = read(DocxCaseExporter().export([ermita]))
    assert not any("[^" in line for line in texts(document))

    raised = [run.text for p in document.paragraphs for run in p.runs if run.font.superscript]
    expected = len(re.findall(r"\[\^\d+\]", ermita.full_text)) + sum(len(re.findall(r"\[\^\d+\]", o.text)) for o in ermita.opinions)
    assert len(raised) == expected and expected > 0


def test_every_footnote_is_listed_with_the_courts_own_words(ermita):
    lines = texts(read(DocxCaseExporter().export([ermita])))
    for footnote in ermita.footnotes:
        assert f"{footnote.number}. {footnote.text}" in lines
    assert "Footnotes" in lines
    assert len(ermita.footnotes) == 42 and len(ermita.opinions) == 1  # read from the page: 42 footnotes + Brion's opinion


def test_the_concurring_opinion_is_included_with_its_own_footnotes(ermita):
    document = read(DocxCaseExporter().export([ermita]))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert any(h.startswith("Concurring opinion, Justice") for h in headings)
    lines = texts(document)
    for footnote in ermita.opinions[0].footnotes:
        assert f"{footnote.number}. {footnote.text}" in lines
    assert len(ermita.opinions[0].footnotes) == 13


def test_the_header_gives_the_number_the_official_page_and_the_source(ermita):
    lines = texts(read(DocxCaseExporter().export([ermita])))
    assert lines[0] == "G.R. No. 180046"
    assert "April 02, 2009" in lines[1]
    first_party = ermita.title.split(",")[0]
    assert not lines[0].startswith(first_party) and any(first_party in line for line in lines[3:])  # the caption is in the decision, where Lawphil prints it
    assert f"Official page: {URL}" in lines
    assert any("Lawphil (Arellano Law Foundation)" in line for line in lines[:5])


def test_a_joint_decision_lists_every_number_it_settles():
    joint = parse("gr_211972_2015.html", "https://lawphil.net/judjuris/juri2015/jul2015/gr_211972_2015.html")
    assert len(joint.all_numbers) > 1
    lines = texts(read(DocxCaseExporter().export([joint])))
    assert any(line.startswith("G.R. Nos. ") and all(n in line for n in joint.all_numbers) for line in lines[:3])


def test_several_cases_get_a_cover_page_and_each_starts_on_a_new_page(ermita):
    other = parse("gr_173931_2009.html", "https://lawphil.net/judjuris/juri2009/jun2009/gr_173931_2009.html")
    document = read(DocxCaseExporter().export([ermita, other]))
    lines = texts(document)
    assert lines[0] == "Full cases"
    bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert len(bullets) == 2 and bullets[0].startswith("G.R. No. 180046: ") and bullets[1].startswith("G.R. No. 173931: ")
    breaks = [r for p in document.paragraphs for r in p.runs if 'w:br' in r._r.xml and 'type="page"' in r._r.xml]
    assert len(breaks) == 2  # before each case
    assert lines.index("G.R. No. 180046") < lines.index("G.R. No. 173931")  # in the order given


def test_a_single_case_has_no_cover_page(ermita):
    assert texts(read(DocxCaseExporter().export([ermita])))[0] == "G.R. No. 180046"


def test_a_very_long_list_of_parties_is_cut_on_the_cover_only(ermita):
    from dataclasses import replace

    long_title = "A. " + ", ".join(f"PETITIONER NUMBER {n}" for n in range(200)) + ", Petitioners, vs. B, Respondent."
    big = replace(ermita, title=long_title)
    document = read(DocxCaseExporter().export([big, ermita]))
    bullet = next(p.text for p in document.paragraphs if p.style.name == "List Bullet")
    assert len(bullet) < 150 and bullet.endswith("…")


def test_footnote_numbers_without_footnote_text_are_said_plainly(ermita):
    from dataclasses import replace

    lines = texts(read(DocxCaseExporter().export([replace(ermita, footnotes=[], opinions=[])])))
    assert any("footnote numbers, but their text could not be read" in line for line in lines)


def test_a_case_without_a_stored_title_uses_its_number(ermita):
    from dataclasses import replace

    untitled = replace(ermita, title=None, opinions=[], footnotes=[], full_text="Short text with no notes.")
    lines = texts(read(DocxCaseExporter().export([untitled])))
    assert lines[0] == "G.R. No. 180046"
    assert "Footnotes" not in lines and not any("could not be read" in line for line in lines)


def test_the_reader_rules_match_the_on_screen_reader():
    blocks = DecisionBlocks().split("Manila\nEN BANC\nG.R. No. 1\nA v. B\nD E C I S I O N\nSANTOS, J.:\nThe Facts\nThe petition is about a long dispute over land.[^1] It began.\nSO ORDERED.")
    kinds = [(b.kind, b.text) for b in blocks]
    assert (BlockKind.TITLE, "DECISION") in kinds
    assert (BlockKind.PONENTE, "SANTOS, J.:") in kinds
    assert (BlockKind.HEADING, "The Facts") in kinds
    assert (BlockKind.PARTIES, "A v. B") in kinds
    assert kinds[-1] == (BlockKind.PARAGRAPH, "SO ORDERED.")
