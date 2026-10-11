"""The case digest Word file: the client's layout, and a level only chooses which sections print."""
import io

import docx
import pytest

from caselens.domain.digest import AnswerSentence
from caselens.domain.digest_v2 import DigestBlock, DigestDraft, Level, Section, sections_for
from caselens.infrastructure.docx_digest_export import DigestHeader, DocxCaseDigestExporter

HEADER = DigestHeader("Marcos v. Manglapus", "G.R. No. 88211, September 15, 1989 (En Banc)", "Constitutional Law", "Cortes, J.")


def block(*texts, heading=None, as_list=False, cites=("P6", "P7")):
    return DigestBlock(tuple(AnswerSentence(t, cites) for t in texts), heading, as_list)


DRAFT = DigestDraft({
    Section.DOCTRINE: (block("The President has residual powers."),),
    Section.FACTS: (block("Marcos was deposed in 1986.", "Aquino became President."),),
    Section.ARGUMENTS_PETITIONERS: (block("They invoked the right to travel.", "They invoked the right to return.", as_list=True),),
    Section.ISSUE: (block("May the President bar the Marcoses from returning? Yes."),),
    Section.RULING: (block("The petition was dismissed."),),
    Section.RATIO: (block("The Court said so.", heading="1. The right to return"),),
    Section.DISSENTS: (block("He argued rights come first.", heading="Cruz, J., dissenting", cites=("O3.2",)),),
    Section.TOPIC: (block("Executive power is more than a list."),),
    Section.WHY: (block("Remember residual powers.", as_list=True),),
    Section.CASE_SUMMARY: (block("Marcos was deposed and wanted to return.", "The Supreme Court dismissed the petition."),),
})
NO_SUMMARY = DigestDraft({s: b for s, b in DRAFT.sections.items() if s is not Section.CASE_SUMMARY})  # a digest written before v5


def read(level, draft=DRAFT, **kwargs):
    document = docx.Document(io.BytesIO(DocxCaseDigestExporter().export(HEADER, draft, level, **kwargs)))
    return [(p.style.name, p.text) for p in document.paragraphs if p.text.strip()]


def headings(level, draft=DRAFT):
    return [text for style, text in read(level, draft) if style.startswith("Heading 1")]


def test_the_header_is_laid_out_like_the_clients_sample():
    lines = [text for _, text in read(Level.FULL)[:4]]
    assert lines == ["CASE DIGEST", "Marcos v. Manglapus", "G.R. No. 88211, September 15, 1989 (En Banc)", "Topic: Constitutional Law  |  Ponente: Cortes, J."]


def test_the_full_digest_has_every_section_in_the_clients_order():
    assert headings(Level.FULL) == ["Case Summary", "Doctrine", "Facts", "Issue", "Ruling", "Ratio Decidendi", "The Dissents (useful for recitation)", "Topic Explained", "Why This Case Matters"]


def test_the_standard_level_stops_after_the_ruling():
    assert headings(Level.STANDARD) == ["Doctrine", "Facts", "Issue", "Ruling"]


def test_the_short_level_is_the_case_summary_then_the_doctrine_like_the_clients_example():
    assert headings(Level.SHORT) == ["Case Summary", "Doctrine"]
    assert ("Normal", "Marcos was deposed and wanted to return. The Supreme Court dismissed the petition.") in read(Level.SHORT)


def test_the_case_summary_starts_the_full_digest_but_not_the_standard_one():
    assert headings(Level.FULL)[0] == "Case Summary" and "Case Summary" not in headings(Level.STANDARD)
    assert headings(Level.FULL, NO_SUMMARY)[:2] == ["Doctrine", "Facts"]  # no summary: no stand-in either, the full digest has them all


def test_a_digest_without_a_case_summary_prints_facts_issue_and_ruling_in_its_place():
    assert headings(Level.SHORT, NO_SUMMARY) == ["Facts", "Issue", "Ruling", "Doctrine"]
    assert sections_for(Level.SHORT, DRAFT) == (Section.CASE_SUMMARY, Section.DOCTRINE)


def test_the_partys_arguments_are_a_subheading_of_the_facts_and_a_list_prints_as_bullets():
    document = read(Level.STANDARD)
    assert ("Heading 2", "Petitioners’ arguments") in document
    assert [text for style, text in document if style == "List Bullet"] == ["They invoked the right to travel.", "They invoked the right to return."]


def test_a_paragraph_block_is_one_paragraph_and_a_block_heading_is_kept():
    document = read(Level.FULL)
    assert ("Normal", "Marcos was deposed in 1986. Aquino became President.") in document
    assert ("Heading 2", "1. The right to return") in document


def test_like_the_clients_sample_there_are_no_paragraph_lines_unless_asked_for():
    assert not any(text.startswith("Based on decision paragraphs") for _, text in read(Level.FULL))
    assert ("Normal", "Based on decision paragraphs 6, 7.") in read(Level.FULL, show_sources=True)


def runs(level=Level.FULL, draft=None):
    document = docx.Document(io.BytesIO(DocxCaseDigestExporter().export(HEADER, draft or DRAFT, level)))
    return [(p.style.name, [(r.text, bool(r.bold)) for r in p.runs]) for p in document.paragraphs if p.text.strip()]


def test_a_key_sentence_is_bold_and_the_rest_of_its_paragraph_is_not():
    key = DigestDraft({Section.RULING: (DigestBlock((AnswerSentence("The petition was DISMISSED.", ("P102",), key=True), AnswerSentence("The Court explained why.", ("P102",)))),)})
    ruling = next(r for style, r in runs(draft=key) if r and r[0][0].startswith("The petition"))
    assert ruling == [("The petition was DISMISSED.", True), (" The Court explained why.", False)]


def test_each_dissent_is_one_item_with_the_justices_name_in_bold_and_the_recitation_points_have_their_lead_in():
    draft = DigestDraft({
        Section.DISSENTS: (DigestBlock((AnswerSentence("Cruz, J.: The government failed to prove a threat.", ("O2.3",)),), as_list=True),),
        Section.WHY: (DigestBlock((AnswerSentence("Residual powers exist.", ("P74",), key=True),), as_list=True),),
    })
    printed = runs(draft=draft)
    assert ("List Bullet", [("Cruz, J.:", True), (" The government failed to prove a threat.", False)]) in printed
    lead = [r for _, r in printed].index([("Remember for recitation:", True)])
    assert printed[lead + 1] == ("List Bullet", [("Residual powers exist.", True)])


def test_it_always_ends_with_the_draft_notice():
    last = read(Level.SHORT)[-1][1]
    assert "draft" in last and "Lawphil" in last


@pytest.mark.parametrize("level", list(Level))
def test_a_section_the_writer_left_empty_is_not_printed(level):
    thin = DigestDraft({Section.FACTS: DRAFT.sections[Section.FACTS]})
    document = docx.Document(io.BytesIO(DocxCaseDigestExporter().export(HEADER, thin, level)))
    assert [p.text for p in document.paragraphs if p.style.name.startswith("Heading 1")] == ["Facts"]


# -- "Case Summary and Doctrine", as the client's reference (Nobleza v. Nuega) prints it ----------------------------------------

NOBLEZA = DigestHeader("Nobleza v. Nuega", "G.R. No. 193038, March 11, 2015 (Third Division)", "Civil Law", "Villarama, Jr.",
                       number="G.R. No. 193038", decided="March 11, 2015", reporter="752 SCRA 602")
REFERENCE = DigestDraft({
    Section.CASE_SUMMARY: (
        block("Shirley and Rogelio Nuega married on September 1, 1990."),
        DigestBlock((AnswerSentence("The Supreme Court denied the petition.", ("P30",), True),)),
    ),
    Section.DOCTRINE: (
        DigestBlock((AnswerSentence("A buyer claiming good faith bears the burden of proving it.", ("P19",), True),
                     AnswerSentence("The ordinary presumption of good faith is not enough.", ("P19",))),
                    "1. Innocent purchaser for value: burden and due diligence"),
        DigestBlock((AnswerSentence("Under Art. 96, a sale without written consent is void.", ("P27",), True),),
                    "2. Sale of community property without written consent is void in its entirety"),
    ),
})


def short_file(header=NOBLEZA, draft=REFERENCE):
    return docx.Document(io.BytesIO(DocxCaseDigestExporter().export(header, draft, Level.SHORT)))


def test_the_short_file_opens_with_the_references_one_line_top():
    first = short_file().paragraphs[0]
    assert first.text == "Nobleza v. Nuega, G.R. No. 193038 (March 11, 2015), 752 SCRA 602"
    assert first.runs[0].text == "Nobleza v. Nuega" and first.runs[0].italic
    assert "CASE DIGEST" not in [p.text for p in short_file().paragraphs]


def test_the_top_line_adds_only_what_is_known():
    from dataclasses import replace

    assert short_file(replace(NOBLEZA, reporter="755 Phil. 656")).paragraphs[0].text == "Nobleza v. Nuega, G.R. No. 193038 (March 11, 2015), 755 Phil. 656"
    assert short_file(replace(NOBLEZA, reporter=None)).paragraphs[0].text == "Nobleza v. Nuega, G.R. No. 193038 (March 11, 2015)"


def test_the_summary_is_plain_and_each_doctrine_is_one_paragraph_led_by_its_label_in_bold():
    paragraphs = [p for p in short_file().paragraphs if p.text.strip()]
    disposition = next(p for p in paragraphs if p.text == "The Supreme Court denied the petition.")
    assert not any(run.bold for run in disposition.runs)  # no bold inside the summary, as in the reference
    first = next(p for p in paragraphs if p.text.startswith("1. "))
    assert first.text == ("1. Innocent purchaser for value: burden and due diligence: A buyer claiming good faith bears the burden of "
                          "proving it. The ordinary presumption of good faith is not enough.")
    assert first.runs[0].bold and first.runs[0].text == "1. Innocent purchaser for value: burden and due diligence: "
    assert not first.runs[1].bold
    assert not any(p.style.name.startswith("Heading 2") for p in paragraphs)  # the label is not a separate heading


def test_a_heading_that_only_repeats_the_section_name_is_not_printed():
    repeated = DigestDraft({Section.DOCTRINE: (DigestBlock((AnswerSentence("The rule.", ("P1",)),), "Doctrine"),)})
    paragraphs = [p.text for p in docx.Document(io.BytesIO(DocxCaseDigestExporter().export(HEADER, repeated, Level.FULL))).paragraphs]
    assert paragraphs.count("Doctrine") == 1 and "The rule." in paragraphs  # the section title once, no "Doctrine:" label
