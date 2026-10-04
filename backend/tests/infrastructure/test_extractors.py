import io
import re
from pathlib import Path

import docx
import pytest

from caselens.domain.errors import DocumentExtractionError, UnsupportedDocumentError
from caselens.domain.services.text_normalizer import TextNormalizer
from caselens.infrastructure.extraction.composite_extractor import CompositeDocumentExtractor
from caselens.infrastructure.extraction.docx_extractor import DocxTextExtractor
from caselens.infrastructure.extraction.pdf_extractor import PdfTextExtractor

FIXTURES = Path(__file__).resolve().parent.parent / "fixtures"
# `pdftotext` counted 3,669 words in the sample before any normalization (Phase 0).
PDFTOTEXT_BASELINE_WORDS = 3669


@pytest.fixture
def extractor() -> CompositeDocumentExtractor:
    return CompositeDocumentExtractor([PdfTextExtractor(), DocxTextExtractor()])


def test_sample_pdf_is_extracted_with_expected_content(extractor):
    text = extractor.read("sample case.pdf", (FIXTURES / "sample_case.pdf").read_bytes())
    assert "Review Center v Ermita" in text
    assert "GR no 180046" in text
    assert "WHY THIS CASE MATTERS" in text


def test_sample_pdf_word_count_close_to_pdftotext_baseline(extractor):
    text = extractor.read("sample case.pdf", (FIXTURES / "sample_case.pdf").read_bytes())
    words = len(text.split())
    assert abs(words - PDFTOTEXT_BASELINE_WORDS) / PDFTOTEXT_BASELINE_WORDS < 0.05, words


def test_normalized_sample_pdf_has_no_junk(extractor):
    raw = extractor.read("sample case.pdf", (FIXTURES / "sample_case.pdf").read_bytes())
    clean = TextNormalizer.default().normalize(raw)
    assert "​" not in clean
    assert not re.search(r"1a(?:w|vv)p", clean)
    assert "Court. The alleged violation" in clean  # glued footnote 21 removed


def test_docx_text_and_tables_are_extracted(extractor):
    document = docx.Document()
    document.add_paragraph("GR no 180046 (April 2, 2009)")
    table = document.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Review Center v Ermita"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extractor.read("reviewer.docx", buffer.getvalue())
    assert "GR no 180046" in text
    assert "Review Center v Ermita" in text


def test_unsupported_extension_is_rejected(extractor):
    with pytest.raises(UnsupportedDocumentError):
        extractor.read("notes.txt", b"hello")


def test_corrupt_pdf_raises_extraction_error(extractor):
    with pytest.raises(DocumentExtractionError):
        extractor.read("broken.pdf", b"not a pdf")
