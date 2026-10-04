"""Upload a reviewer -> it is checked -> a digest starts for every case it cites -> the finished reviewer comes out as a
Word file. Real PostgreSQL and the real parsed GR 180046 decision; the AI is a scripted fake."""
import io

import docx
import pytest
from fastapi.testclient import TestClient

from caselens.composition import Services
from caselens.domain.digest import AnswerSentence
from caselens.main import create_app
from caselens.presentation.dependencies import get_services
from tests.fakes import AlwaysSupported, FakeJobQueue, FixtureCaseSource, ScriptedWriter
from tests.helpers import OFFICIAL_URL, official_html, reviewer_docx

pytestmark = pytest.mark.db


@pytest.fixture
def jobs():
    return FakeJobQueue()


@pytest.fixture
def services(db_session, jobs):
    source = FixtureCaseSource({OFFICIAL_URL: official_html()}, {"180046": [OFFICIAL_URL]})
    writer = ScriptedWriter([AnswerSentence("The Court declared the order void.", ("P132",))])
    return Services(db_session, case_locator=source, case_fetcher=source, jobs=jobs, answer_writer=writer, answer_checker=AlwaysSupported())


@pytest.fixture
def client(services):
    app = create_app()
    app.dependency_overrides[get_services] = lambda: services
    return TestClient(app)


def upload(client, data=None, name="reviewer.docx"):
    response = client.post("/uploads", files={"file": (name, data or reviewer_docx(), "application/octet-stream")})
    return response.json()["id"]


def finish(services, upload_id):
    """What the workers do: check the citations, then write every digest's answers."""
    services.resolve_upload_citations().execute(upload_id)
    for digest in services.get_digest().for_upload(upload_id):
        services.build_case_digest().execute(digest.id)


def test_checking_a_reviewer_starts_a_digest_for_every_case_it_cites(client, services, jobs):
    upload_id = upload(client)
    assert jobs.digest_builds == []  # the case is not stored yet: nothing to digest until the worker finds it
    services.resolve_upload_citations().execute(upload_id)
    digests = client.get(f"/uploads/{upload_id}/digests").json()
    assert len(digests) == 1 and digests[0]["gr_no"] == "180046" and digests[0]["upload_id"] == upload_id
    assert jobs.digest_builds == [(digests[0]["id"], None)]


def test_the_finished_reviewer_places_the_digest_after_the_paragraph_that_cites_the_case(client, services):
    upload_id = upload(client)
    finish(services, upload_id)

    body = client.get(f"/uploads/{upload_id}/document").json()
    assert body["source"] == "docx" and body["all_ready"] is True
    assert [b["index"] for b in body["blocks"] if b["boxes"]] == [2]  # the paragraph that says "GR no 180046"
    box = body["blocks"][2]["boxes"][0]
    assert box["title"] == "Digest 1: Facts, Issue, Ruling and Doctrine" and "G.R. No. 180046 (April 2, 2009)" in box["heading"]


def test_the_word_file_is_the_students_file_plus_the_box_with_the_ai_answer_marked(client, services):
    upload_id = upload(client)
    finish(services, upload_id)

    response = client.get(f"/uploads/{upload_id}/document.docx")
    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="reviewer-with-digests.docx"'
    document = docx.Document(io.BytesIO(response.content))
    cell = "\n".join(p.text for p in document.tables[0].rows[0].cells[0].paragraphs)
    assert "G.R. No. 180046 (April 2, 2009)" in cell  # the Court's record, not the student's 2010
    assert "your reviewer says 2010; the Court's record says 2009" in cell
    assert "WHEREFORE, we GRANT the petition and the petition-in-intervention." in cell  # the ruling, verbatim
    assert "The Court declared the order void.\nDrafted from the decision (decision paragraph 132). Check it." in cell
    assert "Doctrine\n" in cell  # left empty for the student to pick
    assert [p.text for p in document.paragraphs if p.text.startswith("PART NINE")] == ["PART NINE: LEGISLATIVE DEPARTMENT - Article VI 1987 Constitution"]


def test_a_digest_the_student_edited_comes_out_with_their_words(client, services):
    upload_id = upload(client)
    finish(services, upload_id)
    digest_id = client.get(f"/uploads/{upload_id}/digests").json()[0]["id"]
    client.put(f"/digests/{digest_id}/fields/doctrine/passage", json={"first": 109, "last": 109})
    client.put(f"/digests/{digest_id}/fields/topic/text", json={"text": "My own explanation."})

    document = docx.Document(io.BytesIO(client.get(f"/uploads/{upload_id}/document.docx").content))
    cell = "\n".join(p.text for p in document.tables[0].rows[0].cells[0].paragraphs)
    assert "My own explanation." in cell and "The President has no inherent or delegated legislative power" in cell
    assert "The Court declared the order void.\nDrafted from the decision" in cell  # the untouched 'why' answer is still marked as drafted


def test_while_answers_are_still_being_written_the_file_says_so(client, services):
    upload_id = upload(client)
    services.resolve_upload_citations().execute(upload_id)  # digests requested, answers NOT written yet
    assert client.get(f"/uploads/{upload_id}/document").json()["all_ready"] is False
    cell = "\n".join(p.text for p in docx.Document(io.BytesIO(client.get(f"/uploads/{upload_id}/document.docx").content)).tables[0].rows[0].cells[0].paragraphs)
    assert "Still being written." in cell and "Download again in a minute" in cell


def test_a_pdf_is_rebuilt_as_a_word_file(client, services):
    pdf = (__import__("tests.helpers", fromlist=["FIXTURES"]).FIXTURES / "sample_case.pdf").read_bytes()
    upload_id = upload(client, pdf, "sample_case.pdf")
    finish(services, upload_id)

    body = client.get(f"/uploads/{upload_id}/document").json()
    assert body["source"] == "pdf" and sum(len(b["boxes"]) for b in body["blocks"]) + len(body["unplaced"]) == 1
    # the sample is the student's own finished document: the case is only named in their own digest headers, so the new
    # box goes at the end, not into the middle of their boxes
    assert len(body["unplaced"]) == 1 and all(b["boxes"] == [] for b in body["blocks"])
    assert body["own_digests"] == 2  # the student's own "Digest 1" and "Digest 2" are noticed, so the screen can say so
    document = docx.Document(io.BytesIO(client.get(f"/uploads/{upload_id}/document.docx").content))
    assert document.paragraphs[0].text == "sample_case.pdf (with digests)" and len(document.tables) == 1


def test_a_citation_the_court_does_not_have_gets_no_box(client, services):
    text_docx = reviewer_docx(["See GR no 999999 for nothing."])
    upload_id = upload(client, text_docx)
    services.resolve_upload_citations().execute(upload_id)
    assert client.get(f"/uploads/{upload_id}/digests").json() == []
    body = client.get(f"/uploads/{upload_id}/document").json()
    assert all(b["boxes"] == [] for b in body["blocks"]) and body["unplaced"] == []
    assert docx.Document(io.BytesIO(client.get(f"/uploads/{upload_id}/document.docx").content)).tables == []


def test_an_unknown_upload_is_404(client):
    assert client.get("/uploads/999999/document").status_code == 404
    assert client.get("/uploads/999999/document.docx").status_code == 404


def test_the_titles_of_the_reviewer_are_marked_so_the_screen_can_make_them_stand_out(client, services):
    paragraphs = [
        "Sample case",
        "PART NINE: LEGISLATIVE DEPARTMENT - Article VI 1987 Constitution",
        "I. Legislative power Section 1:",
        "Section 1: The legislative power shall be vested in the Congress of the Philippines. See Review Center v Ermita, GR no 180046 (April 2, 2010).",
        "A. General Plenary powers",
        "Plenary Nature: Legislative power is generally considered plenary or absolute.",
    ]
    upload_id = upload(client, reviewer_docx(paragraphs))
    services.resolve_upload_citations().execute(upload_id)
    levels = [b["heading_level"] for b in client.get(f"/uploads/{upload_id}/document").json()["blocks"]]
    assert levels == [1, 1, 2, None, 3, None]


def test_a_plain_reviewer_has_no_digests_of_its_own(client, services):
    upload_id = upload(client)
    services.resolve_upload_citations().execute(upload_id)
    assert client.get(f"/uploads/{upload_id}/document").json()["own_digests"] == 0
