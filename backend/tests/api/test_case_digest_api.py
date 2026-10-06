"""The case digest endpoints on real PostgreSQL, with the writer and the checks scripted (no AI): ask, poll, read, download."""
import io

import docx
import pytest
from fastapi.testclient import TestClient

from caselens.composition import Services
from caselens.domain.digest import AnswerSentence
from caselens.domain.digest_v2 import DigestBlock, DigestDraft, Section
from caselens.infrastructure.db.repositories import SqlCaseRepository
from caselens.main import create_app
from caselens.presentation.dependencies import get_services
from tests.fakes import FakeJobQueue, ScriptedDigestWriter, VerdictChecker
from tests.helpers import parse_digest_case

pytestmark = pytest.mark.db

_WORD = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DRAFT = DigestDraft({
    Section.DOCTRINE: (DigestBlock((AnswerSentence("The President may not make law by executive order.", ("P132",)),)),),
    Section.FACTS: (DigestBlock((AnswerSentence("The PRC reported that exam questions had leaked.", ("P9",)),)),),
    Section.ISSUE: (DigestBlock((AnswerSentence("May the President regulate review centers by order? No.", ("P64", "P65")),)),),
    Section.RULING: (DigestBlock((AnswerSentence("The petition was granted.", ("P132",)),)),),
    Section.RATIO: (DigestBlock((AnswerSentence("Only Congress can give CHED that power.", ("P120",)),), "1. The power is Congress's"),),
})


@pytest.fixture
def jobs():
    return FakeJobQueue()


@pytest.fixture
def services(db_session, jobs):
    return Services(db_session, jobs=jobs, digest_writer=ScriptedDigestWriter(DRAFT), answer_checker=VerdictChecker())


@pytest.fixture
def client(services):
    app = create_app()
    app.dependency_overrides[get_services] = lambda: services
    return TestClient(app)


@pytest.fixture
def case_id(db_session):
    return SqlCaseRepository(db_session).add(parse_digest_case("gr_180046_2009.html")).id


def test_before_anyone_asks_the_state_is_none_and_the_header_is_filled(client, case_id):
    body = client.get(f"/cases/{case_id}/case-digest").json()
    assert body["state"] == "none" and body["sections"] == []
    assert body["header"]["case_name"] == "Review Center Association of the Philippines v. Ermita"
    assert body["header"]["citation"] == "G.R. No. 180046, April 2, 2009 (En Banc)"
    assert set(body["levels"]) == {"short", "standard", "full"}
    assert body["levels"]["short"] == ["doctrine", "facts", "arguments_petitioners", "arguments_respondents"]


def test_asking_returns_202_pending_and_queues_the_job_once(client, case_id, jobs):
    first = client.post(f"/cases/{case_id}/case-digest")
    again = client.post(f"/cases/{case_id}/case-digest")
    assert first.status_code == 202 and first.json()["state"] == "pending" and again.json()["state"] == "pending"
    assert len(jobs.case_digests) == 1  # one digest id, queued once


def test_a_scope_is_its_own_digest_named_in_the_header(client, services, case_id, jobs):
    client.post(f"/cases/{case_id}/case-digest")
    scoped = client.post(f"/cases/{case_id}/case-digest", params={"scope": "Delegation of  legislative power"}).json()
    assert scoped["scope"] == "Delegation of legislative power" and scoped["header"]["topic"] == "Delegation of legislative power"
    assert len(jobs.case_digests) == 2  # the standard digest and the scoped one
    services.build_case_digest_v2().execute(jobs.case_digests[1])
    assert client.get(f"/cases/{case_id}/case-digest", params={"scope": "delegation of legislative power"}).json()["state"] == "ready"
    assert client.get(f"/cases/{case_id}/case-digest").json()["state"] == "pending"
    assert client.get(f"/cases/{case_id}/case-digest.docx", params={"scope": "Delegation of legislative power"}).status_code == 200


def test_after_the_job_the_digest_is_readable_with_its_sources_and_the_levels(client, services, case_id, jobs):
    services.build_case_digest_v2().execute(client.post(f"/cases/{case_id}/case-digest").json() and jobs.case_digests[-1])  # the worker
    body = client.get(f"/cases/{case_id}/case-digest").json()
    assert body["state"] == "ready" and body["written"] == 5 and body["dropped"] == 0
    assert body["current"] is True  # written by today's prompt and checks (an older digest would say False until written again)
    assert [s["key"] for s in body["sections"]] == ["doctrine", "facts", "issue", "ruling", "ratio"]
    ratio = body["sections"][-1]["blocks"][0]
    assert ratio["heading"] == "1. The power is Congress's" and ratio["sentences"][0]["cites"] == ["P120"]


def test_the_word_file_follows_the_level_and_is_refused_until_the_digest_is_ready(client, services, case_id, jobs):
    assert client.get(f"/cases/{case_id}/case-digest.docx").status_code == 409
    services.build_case_digest_v2().execute(client.post(f"/cases/{case_id}/case-digest").json() and jobs.case_digests[-1])

    def headings(level):
        response = client.get(f"/cases/{case_id}/case-digest.docx", params={"level": level})
        assert response.status_code == 200 and response.headers["content-type"] == _WORD
        assert f"GR-180046-{level}-digest.docx" in response.headers["content-disposition"]
        return [p.text for p in docx.Document(io.BytesIO(response.content)).paragraphs if p.style.name == "Heading 1"]

    assert headings("short") == ["Doctrine", "Facts"]
    assert headings("standard") == ["Doctrine", "Facts", "Issue", "Ruling"]
    assert headings("full") == ["Doctrine", "Facts", "Issue", "Ruling", "Ratio Decidendi"]
    assert client.get(f"/cases/{case_id}/case-digest.docx", params={"level": "huge"}).status_code == 422


def test_the_word_file_has_no_paragraph_references_unless_asked_like_the_clients_sample(client, services, case_id, jobs):
    services.build_case_digest_v2().execute(client.post(f"/cases/{case_id}/case-digest").json() and jobs.case_digests[-1])
    plain = client.get(f"/cases/{case_id}/case-digest.docx").content
    with_sources = client.get(f"/cases/{case_id}/case-digest.docx", params={"sources": True}).content
    text = lambda data: " ".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)  # noqa: E731
    assert "Based on decision paragraphs" not in text(plain) and "Based on decision paragraphs" in text(with_sources)

def test_an_unknown_case_is_404(client):
    assert client.get("/cases/999999/case-digest").status_code == 404
    assert client.post("/cases/999999/case-digest").status_code == 404


def test_a_section_edited_in_a_review_shows_there_and_in_its_word_file_but_not_elsewhere(client, services, case_id, jobs):
    review = client.post("/bulk", json={"text": ""}).json()
    other = client.post("/bulk", json={"text": ""}).json()
    digest = client.post(f"/cases/{case_id}/case-digest").json()
    services.build_case_digest_v2().execute(digest["id"])

    url = f"/bulk/{review['id']}/digests/{digest['id']}/sections/facts"
    assert client.put(url, json={"text": "My own facts.\n\n- First event\n- Second event"}).status_code == 204
    mine = client.get(f"/cases/{case_id}/case-digest", params={"batch_id": review["id"]}).json()
    facts = next(s for s in mine["sections"] if s["key"] == "facts")
    assert facts["edited"] is True and facts["text"] == "My own facts.\n\n- First event\n- Second event"
    assert [b["as_list"] for b in facts["blocks"]] == [False, True] and facts["blocks"][0]["sentences"][0]["cites"] == []

    word = client.get(f"/cases/{case_id}/case-digest.docx", params={"batch_id": review["id"]}).content
    assert "My own facts." in [p.text for p in docx.Document(io.BytesIO(word)).paragraphs]

    elsewhere = client.get(f"/cases/{case_id}/case-digest", params={"batch_id": other["id"]}).json()
    shared = client.get(f"/cases/{case_id}/case-digest").json()
    for body in (elsewhere, shared):  # the shared AI digest and other reviews are untouched
        assert next(s for s in body["sections"] if s["key"] == "facts")["edited"] is False

    assert client.delete(url).status_code == 204  # put back the AI version
    back = client.get(f"/cases/{case_id}/case-digest", params={"batch_id": review["id"]}).json()
    assert next(s for s in back["sections"] if s["key"] == "facts")["edited"] is False
    assert client.put(f"/bulk/999999/digests/{digest['id']}/sections/facts", json={"text": "x"}).status_code == 404
    assert client.put(f"/bulk/{review['id']}/digests/{digest['id']}/sections/nonsense", json={"text": "x"}).status_code == 422
