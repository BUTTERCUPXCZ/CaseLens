"""HTTP layer end to end on the real PostgreSQL, with Lawphil and the job queue replaced by fakes."""
import httpx
import pytest
from fastapi.testclient import TestClient

from caselens.composition import Services
from caselens.domain.value_objects import GrNumber
from caselens.infrastructure.lawphil.case_source import LawphilCaseSource
from caselens.infrastructure.lawphil.http_client import ThrottledPageClient
from caselens.infrastructure.lawphil.url_scheme import LawphilUrlScheme
from caselens.main import create_app
from caselens.presentation.dependencies import get_services
from tests.fakes import FakeJobQueue, FixtureCaseSource, InMemoryMonthIndexRepository
from tests.helpers import OFFICIAL_URL, official_html, sample_pdf_bytes

pytestmark = pytest.mark.db


@pytest.fixture
def source() -> FixtureCaseSource:
    return FixtureCaseSource({OFFICIAL_URL: official_html()}, {"180046": [OFFICIAL_URL]})


@pytest.fixture
def jobs() -> FakeJobQueue:
    return FakeJobQueue()


@pytest.fixture
def services(db_session, source, jobs) -> Services:
    return Services(db_session, case_locator=source, case_fetcher=source, jobs=jobs)


@pytest.fixture
def client(services) -> TestClient:
    app = create_app()
    app.dependency_overrides[get_services] = lambda: services
    return TestClient(app)


def upload_sample(client):
    return client.post("/uploads", files={"file": ("sample case.pdf", sample_pdf_bytes(), "application/pdf")})


def test_full_flow_upload_background_fetch_then_mismatch_report(client, services, jobs):
    # 1. Upload: nothing stored yet, so the citation is pending and a job is queued.
    response = upload_sample(client)
    assert response.status_code == 202
    body = response.json()
    assert body["status"] == "processing"
    assert body["citations"][0]["status"] == "pending"
    assert (body["citations"][0]["gr_no"], body["citations"][0]["claimed"]["year"]) == ("180046", 2010)
    assert jobs.resolve_upload_ids == [body["id"]]

    # 2. The worker's job (run directly here).
    services.resolve_upload_citations().execute(body["id"])

    # 3. Poll: the real case was fetched; the year claimed by the reviewer is flagged.
    report = client.get(f"/uploads/{body['id']}").json()
    assert report["status"] == "done"
    (citation,) = report["citations"]
    assert citation["status"] == "mismatch"
    assert citation["mismatches"] == {"year": {"claimed": 2010, "official": 2009}}
    assert citation["unverified"] == ["reporter"]
    assert citation["source_url"] == OFFICIAL_URL
    assert citation["claimed"]["reporter"] == "538 SCRA 428"
    official = citation["case"]  # the record it was checked against, so the UI needs no second request
    assert (official["gr_no"], official["decision_date"], official["ponente"]) == ("180046", "2009-04-02", "CARPIO")
    assert official["title"].startswith("REVIEW CENTER ASSOCIATION") and official["source_url"] == OFFICIAL_URL

    # 4. The stored official case, with footnote deep links back to Lawphil.
    case = client.get(f"/cases/{citation['case_id']}").json()
    assert (case["gr_no"], case["decision_date"], case["ponente"], case["disposition"]) == (
        "180046", "2009-04-02", "CARPIO", "GRANTED",
    )
    assert len(case["footnotes"]) == 42 and len(case["opinions"][0]["footnotes"]) == 13
    assert case["footnotes"][18]["source_url"] == f"{OFFICIAL_URL}#fnt19"
    assert "raw_html" not in case


def test_second_upload_of_the_same_file_is_instant_and_fetches_nothing(client, services, jobs, source):
    first = upload_sample(client).json()
    services.resolve_upload_citations().execute(first["id"])
    fetches_before, jobs_before = len(source.fetch_calls), len(jobs.resolve_upload_ids)

    second = upload_sample(client)

    assert second.status_code == 200  # already complete, not 202
    assert second.json()["status"] == "done"
    assert second.json()["citations"][0]["status"] == "mismatch"
    assert len(source.fetch_calls) == fetches_before
    assert len(jobs.resolve_upload_ids) == jobs_before


def test_search_by_gr_number(client, services, jobs):
    unknown = client.get("/cases", params={"gr_no": "180046", "year": 2009})
    assert unknown.status_code == 202 and unknown.json() == {"status": "pending", "cases": []}
    assert jobs.fetch_calls == [("180046", 2009)]

    no_year = client.get("/cases", params={"gr_no": "180046"})
    assert no_year.status_code == 200 and no_year.json()["status"] == "needs_year"

    services.fetch_case_by_gr_number().execute(GrNumber("180046"), 2009)
    found = client.get("/cases", params={"gr_no": "180046"}).json()
    assert found["status"] == "found" and found["cases"][0]["source_url"] == OFFICIAL_URL


def test_case_insights_endpoint_has_findings_and_footnote_links(client, services):
    stored = services.fetch_case_by_gr_number().execute(GrNumber("180046"), 2009)

    body = client.get(f"/cases/{stored.id}/insights").json()

    assert (body["gr_no"], body["ponente"], body["disposition"], body["division"]) == (
        "180046", "CARPIO", "GRANTED", "EN BANC",
    )
    assert len(body["concurring_justices"]) == 13 and body["concurring_justices"][0] == "REYNATO S. PUNO"
    assert body["opinions"] == [{"kind": "concurring", "author": "BRION"}]
    assert body["ruling"].startswith("WHEREFORE") and "SO ORDERED" not in body["ruling"]
    assert any(s["type"] == "EO" and s["number"] == "566" for s in body["statutes"])
    lpbs = next(c for c in body["cited_cases"] if c["title"] == "LPBS Commercial, Inc. v. Amila")
    assert lpbs["source_url"] == f"{OFFICIAL_URL}#fnt18"  # one click to the citing footnote
    ople = next(c for c in body["cited_cases"] if c["title"] == "Ople v. Torres")
    assert ople["source_url"] == OFFICIAL_URL  # body mention: link to the page itself


def test_trends_say_so_when_there_is_too_little_data(client, services):
    empty = client.get("/insights/trends").json()
    assert (empty["total_cases"], empty["enough_data"]) == (0, False)

    services.fetch_case_by_gr_number().execute(GrNumber("180046"), 2009)
    one = client.get("/insights/trends").json()
    assert (one["total_cases"], one["enough_data"]) == (1, False)
    assert "need at least 5" in one["message"]
    assert one["cases_per_ponente"] == [{"ponente": "CARPIO", "cases": 1}]


def test_every_response_with_official_text_carries_attribution_and_the_disclaimer(client, services):
    stored = services.fetch_case_by_gr_number().execute(GrNumber("180046"), 2009)
    upload = upload_sample(client).json()

    for body in (
        client.get(f"/cases/{stored.id}").json(),
        client.get(f"/cases/{stored.id}/insights").json(),
        upload,
    ):
        assert "lawphil.net" in body["attribution"]["source"]
        assert "no warranty" in body["attribution"]["notice"]
        assert "Supreme Court" in body["attribution"]["notice"]


def docx_bytes(text: str) -> bytes:
    import io

    import docx

    document = docx.Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload_docx(client, text: str):
    return client.post("/uploads", files={"file": ("reviewer.docx", docx_bytes(text), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})


def test_recent_uploads_list_shows_how_each_review_came_out(client, services):
    first = upload_sample(client).json()
    services.resolve_upload_citations().execute(first["id"])

    listing = client.get("/uploads", params={"limit": 5}).json()

    assert [u["filename"] for u in listing] == ["sample case.pdf"]
    summary = listing[0]
    assert (summary["total"], summary["needs_look"], summary["matched"], summary["status"]) == (1, 1, 0, "done")
    assert summary["created_at"] is not None
    assert client.get("/uploads", params={"limit": 0}).status_code == 422
    assert client.get(f"/uploads/{first['id']}").json()["created_at"] is not None


def test_library_lists_pages_and_searches_stored_cases(client, services):
    assert client.get("/library/cases").json() == {"items": [], "total": 0, "limit": 20, "offset": 0}

    services.fetch_case_by_gr_number().execute(GrNumber("180046"), 2009)

    page = client.get("/library/cases").json()
    assert page["total"] == 1
    item = page["items"][0]
    assert (item["gr_no"], item["ponente"], item["disposition"]) == ("180046", "CARPIO", "GRANTED")
    assert "full_text" not in item and "raw_html" not in item  # a list never carries the text

    assert client.get("/library/cases", params={"q": "review center"}).json()["total"] == 1
    assert client.get("/library/cases", params={"q": "G.R. No. 1800"}).json()["total"] == 1
    assert client.get("/library/cases", params={"q": "nobody"}).json()["total"] == 0
    assert client.get("/library/cases", params={"limit": 101}).status_code == 422
    assert client.get("/library/cases", params={"offset": -1}).status_code == 422


def test_pasting_a_lawphil_link_resolves_a_citation_that_had_no_year(client, services, source):
    source.locations = {}  # the lists cannot place this number, so only the pasted link can
    upload = upload_docx(client, "Review Center v Ermita, GR no 180046").json()
    services.resolve_upload_citations().execute(upload["id"])
    stuck = client.get(f"/uploads/{upload['id']}").json()["citations"][0]
    assert stuck["status"] == "not_found" and "no year was given" in stuck["message"]

    response = client.post(
        f"/uploads/{upload['id']}/citations/{stuck['id']}/attach", json={"url": OFFICIAL_URL}
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "done"
    assert body["citations"][0]["status"] == "match"
    assert body["citations"][0]["source_url"] == OFFICIAL_URL


def test_attach_to_an_unknown_citation_is_404(client):
    upload = upload_docx(client, "GR no 180046 (2009)").json()
    assert client.post(f"/uploads/{upload['id']}/citations/999999/attach", json={"url": OFFICIAL_URL}).status_code == 404
    assert client.post("/uploads/999999/citations/1/attach", json={"url": OFFICIAL_URL}).status_code == 404


def test_retry_rechecks_citations_that_could_not_be_checked(client, services, source, jobs):
    upload = upload_docx(client, "GR no 180046 (2009)").json()
    source.down = True
    services.resolve_upload_citations().execute(upload["id"])
    failed = client.get(f"/uploads/{upload['id']}").json()["citations"][0]
    assert failed["status"] == "error"
    jobs_before = len(jobs.resolve_upload_ids)

    response = client.post(f"/uploads/{upload['id']}/retry")

    assert response.status_code == 200
    body = response.json()
    assert (body["status"], body["citations"][0]["status"]) == ("processing", "pending")
    assert len(jobs.resolve_upload_ids) == jobs_before + 1

    source.down = False
    services.resolve_upload_citations().execute(upload["id"])
    assert client.get(f"/uploads/{upload['id']}").json()["citations"][0]["status"] == "match"
    assert client.post("/uploads/999999/retry").status_code == 404


def test_insights_for_unknown_case_is_404(client):
    assert client.get("/cases/999999/insights").status_code == 404


def test_invalid_gr_number_is_422(client):
    assert client.get("/cases", params={"gr_no": "12"}).status_code == 422


def test_unsupported_and_corrupt_uploads(client):
    assert client.post("/uploads", files={"file": ("a.txt", b"hi", "text/plain")}).status_code == 415
    assert client.post("/uploads", files={"file": ("a.pdf", b"nope", "application/pdf")}).status_code == 422


def test_fetch_by_url_stores_the_case_and_never_requests_foreign_urls(db_session, jobs):
    """Uses the real LawphilCaseSource (with a mocked network) so its URL guard is in play."""
    requested: list[str] = []

    def lawphil(request: httpx.Request) -> httpx.Response:
        requested.append(str(request.url))
        return httpx.Response(200, text=official_html())

    pages = ThrottledPageClient(
        httpx.Client(transport=httpx.MockTransport(lawphil)), 0, 0, sleep=lambda s: None
    )
    real = LawphilCaseSource(pages, InMemoryMonthIndexRepository(), LawphilUrlScheme("https://lawphil.net"))
    services = Services(db_session, case_locator=real, case_fetcher=real, jobs=jobs)
    app = create_app()
    app.dependency_overrides[get_services] = lambda: services
    client = TestClient(app)

    ok = client.post("/cases/fetch", json={"url": OFFICIAL_URL})
    assert ok.status_code == 200 and ok.json()["gr_no"] == "180046"

    for bad in ["https://evil.example/judjuris/a.html", "http://lawphil.net/judjuris/a.html", "file:///etc/passwd"]:
        assert client.post("/cases/fetch", json={"url": bad}).status_code == 400
    assert requested == [OFFICIAL_URL]  # nothing else was ever requested


def test_unknown_ids_are_404(client):
    assert client.get("/cases/999999").status_code == 404
    assert client.get("/uploads/999999").status_code == 404


def test_health(client):
    assert client.get("/health").json() == {"status": "ok", "db": "ok", "pg_trgm": True}


def test_deleting_a_review_removes_it_and_its_digests_but_keeps_the_case(client, services):
    upload = upload_docx(client, "GR no 180046 (2009)").json()
    services.resolve_upload_citations().execute(upload["id"])
    case_id = client.get(f"/uploads/{upload['id']}").json()["citations"][0]["case_id"]
    digest = client.post(f"/cases/{case_id}/digest", json={"upload_id": upload["id"], "template": "full"}).json()
    assert client.get(f"/uploads/{upload['id']}/digests").json() != []
    other = upload_docx(client, "GR no 180046 (2009)").json()

    response = client.delete(f"/uploads/{upload['id']}")

    assert response.status_code == 204 and response.content == b""
    assert client.get(f"/uploads/{upload['id']}").status_code == 404
    assert client.get(f"/digests/{digest['id']}").status_code == 404
    assert client.get(f"/cases/{case_id}").status_code == 200  # the Court's text stays in the library
    assert [u["id"] for u in client.get("/uploads").json()] == [other["id"]]  # another review is untouched


def test_deleting_an_unknown_review_is_404(client):
    assert client.delete("/uploads/999999").status_code == 404


_WORD = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _word_lines(response) -> list[str]:
    import io

    import docx

    return [p.text for p in docx.Document(io.BytesIO(response.content)).paragraphs]


def test_one_case_downloads_in_full_as_a_word_file(client, services):
    upload = upload_docx(client, "GR no 180046 (2009)").json()
    services.resolve_upload_citations().execute(upload["id"])
    case_id = client.get(f"/uploads/{upload['id']}").json()["citations"][0]["case_id"]

    response = client.get(f"/cases/{case_id}/document.docx")

    assert response.status_code == 200
    assert response.headers["content-type"] == _WORD
    assert 'filename="GR-180046-full-case.docx"' in response.headers["content-disposition"]
    lines = _word_lines(response)
    assert any("WHEREFORE, we GRANT the petition" in line for line in lines)  # the Court's ruling, verbatim
    assert f"Official page: {OFFICIAL_URL}" in lines
    assert client.get("/cases/999999/document.docx").status_code == 404


def test_all_cases_of_a_review_download_as_one_word_file(client, services):
    upload = upload_docx(client, "GR no 180046 (2009)").json()
    services.resolve_upload_citations().execute(upload["id"])

    response = client.get(f"/uploads/{upload['id']}/cases.docx")

    assert response.status_code == 200 and response.headers["content-type"] == _WORD
    assert 'filename="reviewer-all-cases.docx"' in response.headers["content-disposition"]
    assert any("WHEREFORE, we GRANT the petition" in line for line in _word_lines(response))
    assert client.get("/uploads/999999/cases.docx").status_code == 404


def test_a_review_with_no_found_case_has_nothing_to_download(client):
    upload = upload_docx(client, "GR no 111111 (2001)").json()  # still pending: nothing found yet
    response = client.get(f"/uploads/{upload['id']}/cases.docx")
    assert response.status_code == 400 and "nothing to download" in response.json()["detail"]
